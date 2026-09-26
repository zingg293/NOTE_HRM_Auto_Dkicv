from datetime import datetime, date
from pathlib import Path
import re
from openpyxl import load_workbook
from models import JobRow

START = date(2026, 7, 1)
END = date(2026, 9, 30)

def validate_run_dates(start_text, end_text, deadline_text, jobs_by_file=()):
    """Validate the dates shown in the UI before opening any account."""
    values = {}
    for key, label, raw in (
        ('start', 'Ngày bắt đầu phiếu', start_text),
        ('end', 'Ngày kết thúc phiếu', end_text),
        ('deadline', 'Hạn hoàn thành các dòng', deadline_text),
    ):
        try:
            parsed = datetime.strptime(str(raw).strip(), '%d/%m/%Y').date()
        except ValueError as error:
            raise ValueError(f"{label} phải có dạng DD/MM/YYYY: '{raw}'.") from error
        values[key] = parsed
    if values['start'] > values['end']:
        raise ValueError('Ngày bắt đầu phiếu phải trước hoặc bằng ngày kết thúc phiếu.')
    if not values['start'] <= values['deadline'] <= values['end']:
        raise ValueError('Hạn hoàn thành phải nằm trong khoảng ngày của phiếu.')
    for file_label, jobs in jobs_by_file:
        for job in jobs:
            task_date = datetime.strptime(job.start_date, '%d/%m/%Y').date()
            if not values['start'] <= task_date <= values['deadline']:
                raise ValueError(
                    f"{file_label}, dòng Excel {job.excel_row}: Ngày thực hiện "
                    f"{job.start_date} phải từ ngày bắt đầu phiếu đến hạn hoàn thành."
                )
    return tuple(values[key].strftime('%d/%m/%Y')
                 for key in ('start', 'end', 'deadline'))

def _to_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None

def _fmt_date(value):
    d = _to_date(value)
    return d.strftime("%d/%m/%Y") if d else str(value or "").strip()

def load_jobs(path: str):
    p = Path(path)
    if not p.exists():
        raise ValueError(f"Không tìm thấy file Excel: {p}")

    wb = load_workbook(p, data_only=True, read_only=False)
    if "Kế hoạch công việc cá nhân" not in wb.sheetnames:
        raise ValueError("Excel không có sheet 'Kế hoạch công việc cá nhân'.")

    ws = wb["Kế hoạch công việc cá nhân"]
    jobs = []
    errors = []

    # B=2, D=4, F=6, H=8, I=9
    for r in range(15, ws.max_row + 1):
        b = ws.cell(r, 2).value
        d = ws.cell(r, 4).value
        f = ws.cell(r, 6).value
        h = ws.cell(r, 8).value
        i = ws.cell(r, 9).value

        # Dòng dữ liệu được nhận diện khi có ít nhất ngày/tên/mô tả.
        if all(v in (None, "") for v in (b, d, i)):
            continue

        # Bỏ qua phần chú thích/footer của mẫu nếu chỉ có chữ ở cột B.
        # Một dòng dữ liệu thực phải có ít nhất Tên + Mô tả.
        if (d in (None, "")) and (i in (None, "")):
            continue

        if not b:
            errors.append(f"Dòng Excel {r}: thiếu Ngày thực hiện.")
            continue
        if not d:
            errors.append(f"Dòng Excel {r}: thiếu Tên công việc.")
            continue
        if not i:
            errors.append(f"Dòng Excel {r}: thiếu Mô tả công việc.")
            continue

        dt = _to_date(b)
        if dt is None:
            errors.append(f"Dòng Excel {r}: Ngày thực hiện '{b}' không đúng định dạng.")
            continue
        if not (START <= dt <= END):
            errors.append(
                f"Dòng Excel {r}: Ngày thực hiện {dt.strftime('%d/%m/%Y')} "
                f"ngoài khoảng 01/07/2026–30/09/2026."
            )
            continue

        qty = "" if f is None else str(f)
        jobs.append(JobRow(
            excel_row=r,
            start_date=dt.strftime("%d/%m/%Y"),
            task_name=str(d).strip(),
            quantity=qty,
            source_deadline=_fmt_date(h),
            description=str(i).strip(),
        ))

    if errors:
        raise ValueError("Dữ liệu Excel không hợp lệ:\n- " + "\n- ".join(errors))

    if not jobs:
        raise ValueError("Không tìm thấy dòng công việc nào từ dòng 15 trở xuống.")

    return jobs

def make_clipboard_tsv(jobs):
    # Tương đương copy từ cột B đến I:
    # Ngày | Mã trống | Tên | SPC trống | SL giao | SL quy đổi trống | Hạn nguồn | Mô tả
    # A newline/tab inside an Excel cell must not create another NOTE grid row/column.
    def single_cell(value):
        return re.sub(r'\s+', ' ', str(value or '')).strip()

    rows = []
    for j in jobs:
        rows.append("\t".join([
            single_cell(j.start_date),
            "",
            single_cell(j.task_name),
            "",
            single_cell(j.quantity),
            "",
            single_cell(j.source_deadline),
            single_cell(j.description),
        ]))
    return "\n".join(rows)

def make_vertical_clipboard(values):
    return "\n".join(str(v) for v in values)
